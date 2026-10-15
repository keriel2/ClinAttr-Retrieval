#!/usr/bin/env python3
"""
build_annotation_sheet.py

Two modes.

1) Template (no data needed, runs anywhere):
     python3 build_annotation_sheet.py --template annotation_template.xlsx
   Writes a workbook with bilingual instructions, four worked examples, and an empty
   Items sheet. Hand this to annotators to practise.

2) Real sample (run on the server, where the pair files are):
     python3 build_annotation_sheet.py \
       --source orthopedic=/path/orthopedic_surgery_split.csv:test \
       --source pulmonology_qwen=/path/pulmonology_batch_backup.csv \
       --source psychiatry=/path/psychiatry_triaged_final.csv \
       --source pulmonology_llama=/path/llamagen_triaged.csv \
       --n 100 --overlap 30 --out-dir annotation
   Each --source is  label=path  or  label=path:split_value  (keeps only that split).
   Only rows with approved == 1 (machine-screened) are eligible.
   Writes:
     annotation_A.xlsx    all N items, for the first annotator
     annotation_B.xlsx    the OVERLAP items in a different order, for the second annotator
     annotation_key.csv   hidden metadata (domain, generator, language). Do not share it
                          with annotators.

The sheet shows short excerpts around the changed words, not full passages, and no model
decisions, so annotators are blind to how a pair was produced.
"""
import argparse
import random
import re
import sys
from pathlib import Path

HELP_EN = [
    ("Purpose",
     "Each item has an ORIGINAL passage and an EDITED passage that differ by one change. "
     "The changed words are marked [[like this]]. Read only the excerpts and the two queries."),
    ("How to answer",
     "For each question Q1 to Q4 type 1 (yes), 0 (no) or ? (cannot judge). Use ? only for Q3 and Q4 "
     "when you cannot read the language of the queries. Add a short comment in Note if you are unsure. "
     "Do not discuss items with the other annotator."),
    ("Q1  Plausible",
     "Is the EDITED passage medically plausible, something that could appear in a real clinical text?"),
    ("Q2  One change",
     "Does the edited passage differ from the original ONLY by the marked change, and does the "
     "Attribute label (for example stage or dose) describe that change?"),
    ("Q3  Original fits",
     "Does the ORIGINAL passage answer the ORIGINAL query better than the edited passage does?"),
    ("Q4  Edited fits",
     "Does the EDITED passage answer the EDITED query better than the original passage does, "
     "and do the two queries ask about the same topic?"),
    ("Time", "About one minute per item."),
]
HELP_ZH = [
    ("目的",
     "每个条目有一段“原始”段落和一段“编辑后”段落，两者只有一处不同。改动的词用 [[这样]] 标出。"
     "只需阅读摘录和两个问题。"),
    ("如何回答",
     "对 Q1 到 Q4 分别填 1（是）、0（否）或 ?（无法判断）。只有在看不懂问题所用语言时，才对 Q3 和 Q4 填 ?。"
     "不确定时请在 Note 栏写简短说明。请勿与另一位标注者讨论。"),
    ("Q1 合理性", "编辑后的段落在医学上合理吗？它能出现在真实的临床文本中吗？"),
    ("Q2 只有一处改动",
     "编辑后的段落与原始段落是否只有标记处不同，并且 Attribute（属性）标签（如分期、剂量）与该改动相符？"),
    ("Q3 原始匹配",
     "对于“原始问题”，原始段落是否比编辑后的段落更能回答？"),
    ("Q4 编辑匹配",
     "对于“编辑后的问题”，编辑后的段落是否比原始段落更能回答？并且两个问题的主题相同吗？"),
    ("用时", "每个条目约一分钟。"),
]
ATTR_GLOSS = ("Attributes / 属性: negation 否定; dosage 剂量; frequency 频率; subtype 亚型; "
              "stage 分期; temporal 时间状态; severity 严重程度")

COLUMNS = ["Item", "Lang", "Attribute", "Original query", "Edited query",
           "Original passage (excerpt)", "Edited passage (excerpt)",
           "Q1 plausible", "Q2 one change", "Q3 original fits", "Q4 edited fits", "Note"]
WIDTHS = [8, 6, 11, 38, 38, 55, 55, 9, 9, 9, 9, 28]

EXAMPLES = [
    # item, lang, attribute, q, q', d, d', answers, why
    ("Ex1", "en", "stage",
     "Which treatment is advised for stage II sarcoidosis?",
     "Which treatment is advised for stage IV sarcoidosis?",
     "... corticosteroids are advised for [[stage II]] sarcoidosis with persistent symptoms ...",
     "... corticosteroids are advised for [[stage IV]] sarcoidosis with persistent symptoms ...",
     (1, 1, 1, 1),
     "Only the stage changes, and each query fits its own passage. All yes."),
    ("Ex2", "en", "dosage",
     "What dose of the drug is used in adults?",
     "What dose of the drug is used in adults at high risk?",
     "... adults take [[10 mg once daily]] for maintenance ...",
     "... adults take [[20 mg twice daily]] for maintenance ...",
     (1, 0, 1, 1),
     "Q2 = 0: two things changed (dose and frequency), so one label does not describe the change."),
    ("Ex3", "en", "temporal",
     "How long should the antibiotic course last?",
     "How long should the antibiotic course last for seven years?",
     "... the antibiotic course lasts [[7 days]] ...",
     "... the antibiotic course lasts [[7 years]] ...",
     (0, 1, 1, 1),
     "Q1 = 0: an antibiotic course of seven years is not medically plausible."),
    ("Ex4", "fr", "severity",
     "Quels sont les traitements de l'asthme sévère ?",
     "Quels sont les traitements de l'asthme léger ?",
     "... biologic therapy is used in [[severe]] asthma ...",
     "... biologic therapy is used in [[mild]] asthma ...",
     ("1", "1", "?", "?"),
     "Q3 and Q4 = ? only if you cannot read French. Q1 and Q2 can be judged from the English passages."),
]


def style_header(ws, widths):
    from openpyxl.styles import Alignment, Font, PatternFill
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + c)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3E5C76")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "D2"


def add_validation(ws, first_row, last_row):
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"1,0,?"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Use 1, 0 or ?",
                        error="Type 1 for yes, 0 for no, or ? for cannot judge.")
    ws.add_data_validation(dv)
    dv.add(f"H{first_row}:K{last_row}")


def write_help(wb):
    from openpyxl.styles import Alignment, Font
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110
    ws.append(["Instructions", ""])
    ws["A1"].font = Font(bold=True, size=13)
    for k, v in HELP_EN:
        ws.append([k, v])
    ws.append(["", ""])
    ws.append(["说明", ""])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=13)
    for k, v in HELP_ZH:
        ws.append([k, v])
    ws.append(["", ""])
    ws.append(["", ATTR_GLOSS])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[0].font = Font(bold=True)


def write_examples(wb):
    from openpyxl.styles import Alignment
    ws = wb.create_sheet("Examples")
    ws.append(COLUMNS[:7] + COLUMNS[7:11] + ["Why"])
    for ex in EXAMPLES:
        ws.append(list(ex[:7]) + list(ex[7]) + [ex[8]])
    style_header(ws, WIDTHS[:11] + [60])
    ws.cell(row=len(EXAMPLES) + 3, column=1,
            value="Illustrative examples written for training. They are not taken from the data.")
    for row in ws.iter_rows(min_row=2, max_row=len(EXAMPLES) + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_items(wb, rows, blank_rows=0):
    from openpyxl.styles import Alignment
    ws = wb.create_sheet("Items")
    ws.append(COLUMNS)
    for r in rows:
        ws.append(r + [None] * 5)
    for _ in range(blank_rows):
        ws.append([None] * len(COLUMNS))
    style_header(ws, WIDTHS)
    last = max(ws.max_row, 2)
    add_validation(ws, 2, last)
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def make_template(path):
    from openpyxl import Workbook
    wb = Workbook()
    write_help(wb)
    write_examples(wb)
    write_items(wb, [], blank_rows=5)
    wb.save(path)
    print("saved", path)


def excerpt(passage, span, width=30):
    words = str(passage).split()
    span = str(span or "").strip()
    if span:
        m = re.search(re.escape(span), str(passage), flags=re.IGNORECASE)
        if m:
            before = str(passage)[:m.start()].split()[-width:]
            after = str(passage)[m.end():].split()[:width]
            return "... " + " ".join(before) + " [[" + m.group(0) + "]] " + " ".join(after) + " ..."
    return " ".join(words[:2 * width]) + " ..."


def parse_source(text):
    label, rest = text.split("=", 1)
    split = None
    if rest.count(":") >= 1 and not re.match(r"^[A-Za-z]:[\\/]", rest):
        rest, split = rest.rsplit(":", 1)
    return label, rest, split


def build_real(args):
    import pandas as pd
    from openpyxl import Workbook
    rng = random.Random(args.seed)
    pools = {}
    for src in args.source:
        label, path, split = parse_source(src)
        df = pd.read_csv(path)
        ok = df["approved"].astype(str).str.lower().isin(["1", "true", "yes"])
        df = df[ok].copy()
        if split:
            df = df[df["split"].astype(str) == split]
        if "generator_model" not in df.columns:
            df["generator_model"] = "qwen"
        df["generator_model"] = df["generator_model"].fillna("qwen")
        df["source_label"] = label
        for lang, g in df.groupby("language"):
            pools[(label, lang)] = g
    keys = sorted(pools)
    total = sum(len(pools[k]) for k in keys)
    n = min(args.n, total)
    alloc = {k: 0 for k in keys}
    remaining = n
    while remaining > 0:
        progressed = False
        for k in keys:
            if remaining == 0:
                break
            if alloc[k] < len(pools[k]):
                alloc[k] += 1
                remaining -= 1
                progressed = True
        if not progressed:
            break
    picked = []
    for k in keys:
        g = pools[k]
        idx = list(range(len(g)))
        rng.shuffle(idx)
        picked.append(g.iloc[idx[:alloc[k]]])
    sample = pd.concat(picked).sample(frac=1.0, random_state=args.seed).reset_index(drop=True)
    sample["item"] = [f"P{i + 1:03d}" for i in range(len(sample))]

    def to_row(r):
        return [r["item"], r["language"], r["attribute_type"], r["original_query"],
                r["counterfactual_query"],
                excerpt(r["positive_passage"], r["source_span"]),
                excerpt(r["negative_passage"], r["replacement_span"])]

    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    rows_a = [to_row(r) for _, r in sample.iterrows()]
    wb = Workbook()
    write_help(wb)
    write_examples(wb)
    write_items(wb, rows_a)
    wb.save(out / "annotation_A.xlsx")

    overlap = min(args.overlap, len(rows_a))
    sub = sample.sample(n=overlap, random_state=args.seed + 1)
    rows_b = [to_row(r) for _, r in sub.iterrows()]
    wb = Workbook()
    write_help(wb)
    write_examples(wb)
    write_items(wb, rows_b)
    wb.save(out / "annotation_B.xlsx")

    key = sample[["item", "pair_id", "source_label", "domain", "language",
                  "generator_model", "attribute_type"]].copy()
    key["in_overlap"] = key["item"].isin(set(sub["item"])).astype(int)
    key.to_csv(out / "annotation_key.csv", index=False)
    print(f"sample {len(sample)} items; overlap {overlap}; wrote {out}/annotation_A.xlsx, "
          f"annotation_B.xlsx, annotation_key.csv")
    print(key.groupby(["source_label", "language"]).size().to_string())


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("--template", help="write the practice template workbook and exit")
    ap.add_argument("--source", action="append", default=[], help="label=path[:split]")
    ap.add_argument("--n", type=int, default=100)
    ap.add_argument("--overlap", type=int, default=30)
    ap.add_argument("--out-dir", default="annotation")
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()
    if args.template:
        make_template(args.template)
        return
    if not args.source:
        sys.exit("give at least one --source, or use --template")
    build_real(args)


if __name__ == "__main__":
    main()
